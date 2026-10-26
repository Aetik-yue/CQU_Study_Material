from __future__ import annotations

import re
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "计算机网络学习指南_章节案例版.md"
OUTPUT = ROOT / "计算机网络学习指南_章节案例版.docx"

PAGE_WIDTH_DXA = 12240
PAGE_HEIGHT_DXA = 15840
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120

COLORS = {
    "navy": "183153",
    "blue": "2E74B5",
    "dark_blue": "1F4D78",
    "ink": "1F2937",
    "muted": "667085",
    "soft_blue": "E8EEF5",
    "callout": "EEF6FF",
    "code": "F4F6F9",
    "line": "CBD5E1",
    "white": "FFFFFF",
    "gold": "A66B14",
}


def set_run_font(run, latin="Calibri", east_asia="Microsoft YaHei", size=None,
                 color=None, bold=None, italic=None):
    run.font.name = latin
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        tag = "w:" + side
        node = tc_mar.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def set_row_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = tr_pr.find(qn("w:cantSplit"))
    if cant_split is None:
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)


def set_table_borders(table, color="CBD5E1", size="6", outer=True):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn("w:" + edge))
        if node is None:
            node = OxmlElement("w:" + edge)
            borders.append(node)
        if outer or edge in ("insideH", "insideV"):
            node.set(qn("w:val"), "single")
            node.set(qn("w:sz"), size)
            node.set(qn("w:color"), color)
        else:
            node.set(qn("w:val"), "nil")


def set_table_geometry(table, widths: list[int], indent=TABLE_INDENT_DXA):
    assert sum(widths) == CONTENT_WIDTH_DXA, (widths, sum(widths))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths[min(idx, len(widths) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)
            set_cell_margins(cell)


def column_widths(rows: list[list[str]]) -> list[int]:
    n = max(len(r) for r in rows)
    if n == 1:
        return [CONTENT_WIDTH_DXA]
    if n == 2:
        first = max((len(r[0]) if r else 0) for r in rows)
        if first <= 13:
            return [2700, 6660]
        return [4200, 5160]
    if n == 3:
        return [2300, 3530, 3530]
    if n == 4:
        return [2100, 2420, 2420, 2420]
    if n == 5:
        return [1600, 1940, 1940, 1940, 1940]
    if n == 6:
        return [1320, 1608, 1608, 1608, 1608, 1608]
    base = CONTENT_WIDTH_DXA // n
    widths = [base] * n
    widths[-1] += CONTENT_WIDTH_DXA - sum(widths)
    return widths


def set_paragraph_border(paragraph, side="bottom", color="CBD5E1", size="6", space="1"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    border = p_bdr.find(qn("w:" + side))
    if border is None:
        border = OxmlElement("w:" + side)
        p_bdr.append(border)
    border.set(qn("w:val"), "single")
    border.set(qn("w:sz"), size)
    border.set(qn("w:space"), space)
    border.set(qn("w:color"), color)


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_end)
    set_run_font(run, size=9, color=COLORS["muted"])


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), COLORS["blue"])
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(color)
    r_pr.append(underline)
    new_run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


INLINE_RE = re.compile(r"(\*\*.+?\*\*|`.+?`|\[[^\]]+\]\(https?://[^)]+\))")


def add_inline(paragraph, text, base_size=None, base_color=None):
    pos = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos:match.start()])
            set_run_font(run, size=base_size, color=base_color)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=base_size, color=base_color, bold=True)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, latin="Consolas", east_asia="Microsoft YaHei", size=(base_size or 11) - 0.5,
                         color=COLORS["dark_blue"])
            rpr = run._element.get_or_add_rPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), "EEF2F6")
            rpr.append(shd)
        else:
            link_match = re.match(r"\[([^\]]+)\]\((https?://[^)]+)\)", token)
            add_hyperlink(paragraph, link_match.group(1), link_match.group(2))
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_run_font(run, size=base_size, color=base_color)


def setup_styles(doc):
    styles = doc.styles

    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(COLORS["ink"])
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.widow_control = True

    for name, size, color, before, after in (
        ("Heading 1", 16, COLORS["blue"], 18, 10),
        ("Heading 2", 13, COLORS["blue"], 14, 7),
        ("Heading 3", 12, COLORS["dark_blue"], 10, 5),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.05
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    if "Guide Caption" not in styles:
        cap = styles.add_style("Guide Caption", WD_STYLE_TYPE.PARAGRAPH)
    else:
        cap = styles["Guide Caption"]
    cap.font.name = "Calibri"
    cap._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    cap.font.size = Pt(9)
    cap.font.color.rgb = RGBColor.from_string(COLORS["muted"])
    cap.paragraph_format.space_before = Pt(4)
    cap.paragraph_format.space_after = Pt(4)
    cap.paragraph_format.keep_with_next = True


def add_abstract_numbering(doc, fmt, text, abstract_id):
    numbering = doc.part.numbering_part.element
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    nsid = OxmlElement("w:nsid")
    nsid.set(qn("w:val"), f"A{abstract_id:07X}"[-8:])
    abstract.append(nsid)
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), fmt)
    lvl.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), text)
    lvl.append(lvl_text)
    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")
    lvl.append(lvl_jc)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    p_pr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "300")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.append(spacing)
    lvl.append(p_pr)
    r_pr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), "Calibri")
    rfonts.set(qn("w:hAnsi"), "Calibri")
    rfonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r_pr.append(rfonts)
    lvl.append(r_pr)
    abstract.append(lvl)
    numbering.append(abstract)


def next_num_id(doc):
    numbering = doc.part.numbering_part.element
    ids = [int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))]
    return max(ids, default=0) + 1


def new_num_instance(doc, abstract_id):
    numbering = doc.part.numbering_part.element
    num_id = next_num_id(doc)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_num_id = OxmlElement("w:abstractNumId")
    abstract_num_id.set(qn("w:val"), str(abstract_id))
    num.append(abstract_num_id)
    lvl_override = OxmlElement("w:lvlOverride")
    lvl_override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), "1")
    lvl_override.append(start_override)
    num.append(lvl_override)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_id_el)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.25


def setup_page(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("计算机网络学习指南")
    set_run_font(run, size=8.5, color=COLORS["muted"], bold=True)
    run = p.add_run("    COMPUTER NETWORKS · STUDY GUIDE")
    set_run_font(run, size=8.5, color=COLORS["muted"])

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("章节案例版  ·  ")
    set_run_font(run, size=9, color=COLORS["muted"])
    add_page_field(p)


def add_cover(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(92)
    p.paragraph_format.space_after = Pt(22)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("COMPUTER NETWORKS · CHAPTER-BASED FIELD GUIDE")
    set_run_font(run, size=10, color=COLORS["gold"], bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run("计算机网络学习指南")
    set_run_font(run, size=30, color=COLORS["navy"], bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run("从一次网页访问到整个互联网")
    set_run_font(run, size=16, color=COLORS["dark_blue"])

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(54)
    run = p.add_run("知识点 · 典型案例 · 抓包实验 · 高频题型 · 14 周路线")
    set_run_font(run, size=11, color=COLORS["muted"])

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run("零基础入门  |  课程考试  |  408 复习  |  网络编程实践")
    set_run_font(run, size=10.5, color=COLORS["navy"], bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(58)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run("章节案例版")
    set_run_font(run, size=11, color=COLORS["navy"], bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("以协议要解决的真实问题为线索，把五层模型连成一条完整旅程")
    set_run_font(run, size=9.5, color=COLORS["muted"], italic=True)
    doc.add_page_break()


CHAPTER_SUMMARIES = [
    ("第一章  网络全景", "分层、交换、封装，以及发送/传播/排队时延"),
    ("第二章  应用层", "HTTP、DNS、缓存、CDN、邮件与 FTP"),
    ("第三章  传输层", "端口、UDP、TCP 字节流、可靠性和拥塞控制"),
    ("第四章  网络层", "IPv4、CIDR、分片、NAT、DHCP、ICMP、IPv6"),
    ("第五章  路由", "DV、LS、Dijkstra、RIP、OSPF 与 BGP"),
    ("第六章  链路层", "以太网、交换机、ARP、CRC 与 VLAN"),
    ("第七章  物理与无线", "信号、容量、复用、Wi-Fi 与 CSMA/CA"),
    ("第八章  网络安全", "密码机制、证书链、TLS 与常见攻击"),
    ("第九章  综合案例", "网页、会议、游戏、数据库和分层排障"),
    ("第十章  Wireshark", "四组实验、过滤器与证据式实验报告"),
    ("第十一章  高频题型", "子网、时延、分片、CRC、窗口、路由"),
    ("第十二章  学习路线", "14 周安排、费曼笔记、项目与毕业标准"),
]


def add_reading_map(doc):
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.space_before = Pt(0)
    p.add_run("阅读地图")
    intro = doc.add_paragraph()
    add_inline(intro, "主线采用自顶向下的学习顺序；需要考试复习时，可把第十一章当作题型索引，把附录 B 当作口头自测答案。")

    rows = [["章节", "你会解决的问题"]] + [list(x) for x in CHAPTER_SUMMARIES]
    table = doc.add_table(rows=len(rows), cols=2)
    set_table_geometry(table, [2700, 6660])
    set_table_borders(table)
    set_repeat_table_header(table.rows[0])
    for r_idx, row in enumerate(rows):
        set_row_cant_split(table.rows[r_idx])
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.05
            run = p.add_run(value)
            set_run_font(run, size=9.2 if r_idx else 9.5,
                         color=COLORS["navy"] if r_idx == 0 else COLORS["ink"],
                         bold=(r_idx == 0 or c_idx == 0))
            if r_idx == 0:
                set_cell_shading(cell, COLORS["soft_blue"])
    doc.add_page_break()


def add_callout(doc, lines):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [CONTENT_WIDTH_DXA])
    set_row_cant_split(table.rows[0])
    set_table_borders(table, color="A8C7E8", size="8")
    cell = table.cell(0, 0)
    set_cell_shading(cell, COLORS["callout"])
    set_cell_margins(cell, top=120, start=180, bottom=120, end=180)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.18
    add_inline(p, " ".join(lines), base_size=10.2, base_color=COLORS["dark_blue"])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_code_block(doc, text, language=None):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [CONTENT_WIDTH_DXA])
    set_row_cant_split(table.rows[0])
    set_table_borders(table, color="D7DEE8", size="6")
    cell = table.cell(0, 0)
    set_cell_shading(cell, COLORS["code"])
    set_cell_margins(cell, top=110, start=160, bottom=110, end=160)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    lines = text.rstrip("\n").split("\n")
    for idx, line in enumerate(lines):
        run = p.add_run(line)
        set_run_font(run, latin="Consolas", east_asia="Microsoft YaHei", size=8.7, color="263548")
        if idx < len(lines) - 1:
            run.add_break()
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_markdown_table(doc, raw_lines):
    parsed = []
    for line in raw_lines:
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        parsed.append(cells)
    if len(parsed) >= 2 and all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) for c in parsed[1]):
        parsed.pop(1)
    ncols = max(len(r) for r in parsed)
    parsed = [r + [""] * (ncols - len(r)) for r in parsed]
    table = doc.add_table(rows=len(parsed), cols=ncols)
    widths = column_widths(parsed)
    set_table_geometry(table, widths)
    set_table_borders(table)
    if parsed:
        set_repeat_table_header(table.rows[0])
    for r_idx, row in enumerate(parsed):
        set_row_cant_split(table.rows[r_idx])
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            align_center = c_idx > 0 and (len(value) <= 13 or re.fullmatch(r"[\d\s./~～:+×^=<>\-]+", value or "") is not None)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if align_center else WD_ALIGN_PARAGRAPH.LEFT
            add_inline(p, value, base_size=8.6 if ncols >= 5 else 9.1,
                       base_color=COLORS["navy"] if r_idx == 0 else COLORS["ink"])
            if r_idx == 0:
                set_cell_shading(cell, COLORS["soft_blue"])
                for run in p.runs:
                    run.bold = True
    after = doc.add_paragraph(style="Guide Caption")
    after.add_run("")


def add_divider(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    set_paragraph_border(p, color=COLORS["line"], size="6")


def parse_markdown(doc, text):
    lines = text.splitlines()
    start = next(i for i, line in enumerate(lines) if line.startswith("## 导读"))
    lines = lines[start:]
    bullet_abstract = 101
    decimal_abstract = 102
    add_abstract_numbering(doc, "bullet", "•", bullet_abstract)
    add_abstract_numbering(doc, "decimal", "%1.", decimal_abstract)
    bullet_num = new_num_instance(doc, bullet_abstract)
    current_decimal_num = None
    in_decimal_block = False

    i = 0
    paragraph_buffer = []

    def flush_paragraph():
        nonlocal paragraph_buffer
        if paragraph_buffer:
            p = doc.add_paragraph()
            add_inline(p, " ".join(s.strip() for s in paragraph_buffer))
            paragraph_buffer = []

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            flush_paragraph()
            j = i + 1
            while j < len(lines) and not lines[j].strip():
                j += 1
            next_is_decimal = j < len(lines) and re.match(r"^\d+\.\s+", lines[j].strip()) is not None
            if not next_is_decimal:
                in_decimal_block = False
                current_decimal_num = None
            i += 1
            continue

        if stripped.startswith("```"):
            flush_paragraph()
            language = stripped[3:].strip()
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            add_code_block(doc, "\n".join(code_lines), language)
            i += 1
            continue

        if stripped.startswith("|"):
            flush_paragraph()
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i])
                i += 1
            add_markdown_table(doc, table_lines)
            in_decimal_block = False
            continue

        if stripped.startswith(">"):
            flush_paragraph()
            quote_lines = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                quote_lines.append(lines[i].strip()[1:].strip())
                i += 1
            add_callout(doc, quote_lines)
            in_decimal_block = False
            continue

        if stripped == "---":
            flush_paragraph()
            add_divider(doc)
            i += 1
            continue

        heading = re.match(r"^(#{2,4})\s+(.+)$", stripped)
        if heading:
            flush_paragraph()
            level = len(heading.group(1)) - 1
            title = heading.group(2)
            if level == 1:
                p = doc.add_paragraph(style="Heading 1")
                # Heading-level pagination avoids a stray blank page when the
                # preceding chapter ends exactly at a page boundary.
                if title != "导读：怎样把计算机网络真正学会":
                    p.paragraph_format.page_break_before = True
            elif level == 2:
                p = doc.add_paragraph(style="Heading 2")
            else:
                p = doc.add_paragraph(style="Heading 3")
            add_inline(p, title)
            i += 1
            in_decimal_block = False
            continue

        bullet = re.match(r"^-\s+(.+)$", stripped)
        if bullet:
            flush_paragraph()
            p = doc.add_paragraph()
            apply_numbering(p, bullet_num)
            add_inline(p, bullet.group(1))
            i += 1
            in_decimal_block = False
            continue

        number = re.match(r"^\d+\.\s+(.+)$", stripped)
        if number:
            flush_paragraph()
            if not in_decimal_block or current_decimal_num is None:
                current_decimal_num = new_num_instance(doc, decimal_abstract)
                in_decimal_block = True
            p = doc.add_paragraph()
            apply_numbering(p, current_decimal_num)
            add_inline(p, number.group(1))
            i += 1
            continue

        paragraph_buffer.append(stripped)
        i += 1

    flush_paragraph()


def structural_audit(doc):
    assert len(doc.sections) == 1
    section = doc.sections[0]
    assert section.page_width.twips == PAGE_WIDTH_DXA
    assert section.page_height.twips == PAGE_HEIGHT_DXA
    assert section.left_margin.twips == 1440
    assert section.right_margin.twips == 1440
    assert section.top_margin.twips == 1440
    assert section.bottom_margin.twips == 1440
    assert len(doc.tables) > 20
    assert sum(1 for p in doc.paragraphs if p.style.name == "Heading 1") >= 15
    for table in doc.tables:
        tbl_pr = table._tbl.tblPr
        assert tbl_pr.find(qn("w:tblW")) is not None
        assert tbl_pr.find(qn("w:tblLayout")) is not None
        assert table._tbl.tblGrid is not None


def main():
    text = SOURCE.read_text(encoding="utf-8")
    doc = Document()
    setup_styles(doc)
    setup_page(doc)
    doc.core_properties.title = "计算机网络学习指南：从一次网页访问到整个互联网"
    doc.core_properties.subject = "章节知识点、典型案例、Wireshark 实验与复习路线"
    doc.core_properties.keywords = "计算机网络, TCP/IP, HTTP, DNS, TCP, IPv4, Wireshark, TLS"
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    add_cover(doc)
    add_reading_map(doc)
    parse_markdown(doc, text)
    structural_audit(doc)
    doc.save(OUTPUT)
    print(f"Created {OUTPUT}")
    print(f"Paragraphs: {len(doc.paragraphs)}")
    print(f"Tables: {len(doc.tables)}")


if __name__ == "__main__":
    main()
